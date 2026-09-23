"""Tests for GET /attendance/export.xlsx — the monthly payroll timesheet.

Run with the standard library only, alongside the ADMS suite:

    python -m unittest discover -s tests -v

What matters here is not that a file comes back but *what is in it*. Five
things are worth guarding:

  1. The sheet is ZKTime.Net's report, column for column. HR reconciles it
     against payroll and against older files from the terminal's own
     software, so the fourteen Vietnamese headers and their order are a
     contract, not a preference.
  2. Every number in it is a punch scored against the shift. `ScoringTests`
     holds rows lifted from the operator's real August 2026 export, each one
     pinning down a rule the ordinary rows do not.
  3. The export follows the filter, not the page. The table on screen shows 50
     rows; an operator who picks a month and gets 50 of their 900 punches has
     been handed a wrong timesheet with no error anywhere.
  4. The day is paired by *first* and *last* punch. It must not go by the
     device's status field: on the live installation 96,002 of 96,075 records
     carry status=1, so selecting by status yields a sheet of check-outs with
     no check-ins.
  5. An export larger than the configured ceiling is refused with a message,
     not built.

Everything runs against a throwaway in-memory SQLite database.
"""

import os
import unittest
from datetime import date, datetime, time
from io import BytesIO

# Set before importing anything from `app` — see tests/test_adms.py for why.
os.environ.setdefault("APP_ENV", "development")
os.environ.setdefault("SECRET_KEY", "x" * 48)
os.environ.setdefault("DEFAULT_DEVICE_TIMEZONE", "Asia/Dubai")

from fastapi import FastAPI                                    # noqa: E402
from fastapi.testclient import TestClient                      # noqa: E402
from openpyxl import load_workbook                             # noqa: E402
from sqlalchemy import create_engine                           # noqa: E402
from sqlalchemy.orm import sessionmaker                        # noqa: E402
from sqlalchemy.pool import StaticPool                         # noqa: E402

from app import config                                         # noqa: E402
from app.database import Base, get_db                          # noqa: E402
from app.deps import require_auth                              # noqa: E402
from app.models import (                                       # noqa: E402
    AttendanceLog, AuditLog, Device, DeviceEmployee, Employee, User,
)
from app.routers import attendance as attendance_router        # noqa: E402

MAIN_SN = "EXPORTDEV0001"
OTHER_SN = "EXPORTDEV0002"

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# The status every punch on the operator's live installation carries, because
# nobody presses the mode key on the terminal. Used throughout so these tests
# exercise the data as it really arrives, not an idealised in/out alternation.
UNPRESSED = 1


class ExportTestCase(unittest.TestCase):
    """A fresh in-memory database, a signed-in operator and a client."""

    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(bind=self.engine)
        self.Session = sessionmaker(bind=self.engine, autocommit=False, autoflush=False)

        app = FastAPI()
        app.include_router(attendance_router.router)

        def _override_get_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        # Being signed in is a precondition of this endpoint, not its subject.
        operator = User(id=1, username="tester", role="viewer", password_hash="x")
        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[require_auth] = lambda: operator
        self.client = TestClient(app, client=("203.0.113.10", 40000))

        db = self.Session()
        try:
            db.add(Device(serial_number=MAIN_SN, ip_address="203.0.113.10", port=4370,
                          name="Main Door", status="approved", timezone="Asia/Dubai"))
            db.add(Device(serial_number=OTHER_SN, ip_address="203.0.113.11", port=4370,
                          name="Warehouse", status="approved", timezone="Europe/London"))
            db.add(Employee(user_id="1001", name="Nguyen Van A"))
            db.add(Employee(user_id="1002", name="Tran Thi B"))
            db.commit()
        finally:
            db.close()

        self.seed()

    def tearDown(self):
        self.client.close()
        Base.metadata.drop_all(bind=self.engine)
        self.engine.dispose()

    # -- helpers ---------------------------------------------------------

    def seed(self):
        """Per-suite fixture data. Overridden below."""

    def punch(self, moment, user_id="1001", sn=MAIN_SN, status=UNPRESSED,
              punch=1, zone="Asia/Dubai", source="adms_push"):
        db = self.Session()
        try:
            db.add(AttendanceLog(device_sn=sn, user_id=user_id, timestamp=moment,
                                 status=status, punch=punch, source=source,
                                 timezone=zone))
            db.commit()
        finally:
            db.close()

    def export(self, **params):
        return self.client.get("/attendance/export.xlsx", params=params)

    def sheet(self, response):
        wb = load_workbook(BytesIO(response.content))
        return wb.worksheets[0]

    def rows(self, response):
        """Data rows only, as lists of values — the header is row 1."""
        ws = self.sheet(response)
        return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


# The report's fourteen columns, by position. Named so a test reads as the
# sheet reads rather than as a string of magic indexes.
PIN, NAME, DAY, TIMETABLE = 0, 1, 2, 3
ACTUAL, REQUIRED = 4, 5
OT1, OT2, OT3 = 6, 7, 8
LATE, EARLY, ABSENT = 9, 10, 11
CHECK_IN, CHECK_OUT = 12, 13

# The shift these tests are written against, and the one the sample file was
# scored by: 08:30–17:30 with 12:00–13:00 unpaid, Monday to Friday.
FULL_DAY = "8:00"
NOTHING = "0:00"


class ScoringTests(unittest.TestCase):
    """`_score_day` against rows lifted from the real ZKTime.Net export.

    These eight are not invented: each is a row of the operator's August 2026
    file, chosen because it pins down a rule that the ordinary rows do not.
    The whole file (651 rows) reproduces exactly, and these are the ones worth
    keeping in front of the next person to touch the formulas.
    """

    def score(self, came, went, work_day=True):
        day = date(2026, 8, 3)                     # a Monday
        first = datetime.combine(day, time(*came)) if came else None
        last = datetime.combine(day, time(*went)) if went else None
        punches = (1 if first else 0) + (1 if last else 0)
        scored = attendance_router._score_day(first, last, punches, work_day)
        return {k: attendance_router._hm(v) for k, v in scored.items() if k != "paired"}

    def test_an_ordinary_day_loses_the_lunch_hour(self):
        # 08:26 → 17:40 is 9:14 on the clock and 8:14 of work.
        self.assertEqual(
            self.score((8, 26), (17, 40)),
            {"actual": "8:14", "required": FULL_DAY,
             "late": NOTHING, "early": NOTHING, "absent": NOTHING},
        )

    def test_early_arrival_is_worked_not_rounded_to_the_shift(self):
        """In at 08:03 for an 08:30 shift: those 27 minutes are paid, and the
        day is not late. ZKTime.Net paid them, so this does too."""
        scored = self.score((8, 3), (17, 26))
        self.assertEqual(scored["actual"], "8:23")
        self.assertEqual(scored["late"], NOTHING)
        self.assertEqual(scored["early"], "0:04")   # 17:26 is four minutes short

    def test_lateness_is_measured_from_the_shift_start(self):
        self.assertEqual(self.score((8, 43), (17, 39))["late"], "0:13")

    def test_leaving_before_the_break_keeps_the_whole_morning(self):
        """Out at 12:01: the lunch hour is not deducted from work that ended
        before it did, and the 5:29 still owed is 4:30 of *paid* time."""
        scored = self.score((8, 20), (12, 1))
        self.assertEqual(scored["actual"], "3:41")
        self.assertEqual(scored["early"], "4:30")

    def test_arriving_during_the_break_is_not_charged_for_it_twice(self):
        """In at 12:04: the lunch hour was already gone, so none of it comes
        off the afternoon — but the lateness is 3:30, not 3:34, because the
        four minutes of lunch missed were never owed."""
        scored = self.score((12, 4), (17, 51))
        self.assertEqual(scored["actual"], "5:47")
        self.assertEqual(scored["late"], "3:30")

    def test_arriving_before_the_break_and_staying_loses_it(self):
        scored = self.score((11, 24), (17, 51))
        self.assertEqual(scored["actual"], "5:27")
        self.assertEqual(scored["late"], "2:54")

    def test_a_half_day_is_late_and_early_at_once(self):
        scored = self.score((9, 0), (12, 6))
        self.assertEqual(scored["actual"], "3:06")
        self.assertEqual(scored["late"], "0:30")
        self.assertEqual(scored["early"], "4:30")

    def test_one_punch_is_a_whole_day_absent(self):
        """Not a zero-hour day and not a late one: half a record. Charging
        lateness on a day already counted absent would take the same hour
        off twice, and inventing a check-out would pay hours nobody worked."""
        self.assertEqual(
            self.score((8, 36), None),
            {"actual": NOTHING, "required": FULL_DAY,
             "late": NOTHING, "early": NOTHING, "absent": FULL_DAY},
        )

    def test_two_punches_in_the_same_minute_are_still_half_a_record(self):
        """Somebody touching the terminal twice on the way in has not worked
        a day of no hours — they have punched once."""
        day = date(2026, 8, 3)
        first = datetime.combine(day, time(8, 26, 0))
        last = datetime.combine(day, time(8, 26, 40))
        scored = attendance_router._score_day(first, last, 2, True)
        self.assertFalse(scored["paired"])
        self.assertEqual(attendance_router._hm(scored["absent"]), FULL_DAY)

    def test_no_punch_at_all_is_absent(self):
        self.assertEqual(self.score(None, None)["absent"], FULL_DAY)

    def test_a_rest_day_owes_nothing_and_counts_no_absence(self):
        scored = self.score(None, None, work_day=False)
        self.assertEqual(scored["required"], NOTHING)
        self.assertEqual(scored["absent"], NOTHING)

    def test_work_on_a_rest_day_is_still_counted(self):
        """Nothing is owed on a Sunday, but somebody who came in worked."""
        scored = self.score((8, 26), (17, 40), work_day=False)
        self.assertEqual(scored["actual"], "8:14")
        self.assertEqual(scored["required"], NOTHING)
        self.assertEqual(scored["late"], NOTHING)

    def test_durations_are_not_clock_times(self):
        """H:MM, unpadded and uncapped — a duration, not a time of day."""
        self.assertEqual(attendance_router._hm(0), "0:00")
        self.assertEqual(attendance_router._hm(9), "0:09")
        self.assertEqual(attendance_router._hm(494), "8:14")
        self.assertEqual(attendance_router._hm(1500), "25:00")
        self.assertEqual(attendance_router._hm(-5), "0:00")


class DailyExportTests(ExportTestCase):
    """The monthly timesheet, in ZKTime.Net's own layout.

    The week under test is Tuesday 2026-09-01 to Sunday 2026-09-06, so one
    sheet holds worked days, a broken record, an absence and a weekend.
    """

    # The whole week, for both people on the roster.
    WEEK = {"from_date": "2026-09-01T00:00:00", "to_date": "2026-09-06T23:59:59"}

    def seed(self):
        # 1001: a normal Tuesday, a late Wednesday, a Thursday they forgot to
        # punch out of, and a Friday they never came in at all. Every punch
        # carries the same status, exactly as the live terminals send them —
        # the pairing must come from the clock, not from the status field.
        self.punch(datetime(2026, 9, 1, 8, 26, 0))
        self.punch(datetime(2026, 9, 1, 12, 0, 30))      # a midday trip out
        self.punch(datetime(2026, 9, 1, 17, 40, 0))
        self.punch(datetime(2026, 9, 2, 8, 43, 0))
        self.punch(datetime(2026, 9, 2, 17, 39, 0))
        self.punch(datetime(2026, 9, 3, 8, 36, 0))

        # 1002: out before lunch on the Wednesday, in during lunch on the
        # Thursday — the two rows the break rules turn on.
        self.punch(datetime(2026, 9, 2, 8, 20, 0), user_id="1002", sn=OTHER_SN, zone=None)
        self.punch(datetime(2026, 9, 2, 12, 1, 0), user_id="1002", sn=OTHER_SN, zone=None)
        self.punch(datetime(2026, 9, 3, 12, 4, 0), user_id="1002", sn=OTHER_SN, zone=None)
        self.punch(datetime(2026, 9, 3, 17, 51, 0), user_id="1002", sn=OTHER_SN, zone=None)

    def week(self, **params):
        return self.rows(self.export(**{**self.WEEK, **params}))

    def day(self, rows, pin, day):
        """The one row for this person on this date."""
        match = [r for r in rows if r[PIN] == pin and r[DAY] == day]
        self.assertEqual(len(match), 1, f"{pin} on {day}: {len(match)} rows")
        return match[0]

    def test_the_bare_endpoint_is_the_timesheet(self):
        """No parameters at all — what the Export button sends when nothing
        is filtered. There is one kind of export and this is it."""
        res = self.client.get("/attendance/export.xlsx")
        self.assertEqual(res.status_code, 200)
        self.assertIn("attendance-timesheet", res.headers["content-disposition"])

    def test_returns_an_xlsx_with_a_named_attachment(self):
        res = self.export(**self.WEEK)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers["content-type"], XLSX_MEDIA_TYPE)
        self.assertIn("attachment;", res.headers["content-disposition"])
        self.assertIn("attendance-timesheet", res.headers["content-disposition"])
        # A real workbook, not an error page that happens to have the type.
        self.assertTrue(res.content.startswith(b"PK"))

    def test_header_row_is_zktimes_own_columns(self):
        """Verbatim, trailing spaces included. HR reconciles this sheet
        against older files from the terminal's software and against payroll:
        the wording and the order are a contract, not a preference."""
        header = [c.value for c in self.sheet(self.export(**self.WEEK))[1]]
        self.assertEqual(header, [
            "ID nhân viên ", "Họ và tên", "Ngày", "Bảng thời gian",
            "Actual Work", "Require Work",
            "Tăng ca loại 1 ", "Tăng ca loại 2 ", "Tăng ca loại 3 ",
            "Vô trễ", "Ra sớm", "Vắng mặt ", "Check-In", "Check-Out",
        ])

    def test_a_row_for_every_person_on_every_day(self):
        """A timesheet is a grid, not a list of punches: six days for two
        people is twelve rows, whoever did or did not turn up."""
        rows = self.week()
        self.assertEqual(len(rows), 12)
        self.assertEqual([r[PIN] for r in rows], ["1001"] * 6 + ["1002"] * 6)
        self.assertEqual(
            [r[DAY] for r in rows[:6]],
            ["01/09/2026", "02/09/2026", "03/09/2026",
             "04/09/2026", "05/09/2026", "06/09/2026"],
        )
        self.assertEqual(rows[0][NAME], "Nguyen Van A")

    def test_a_worked_day_is_scored_against_the_shift(self):
        row = self.day(self.week(), "1001", "01/09/2026")
        self.assertEqual(row[TIMETABLE], "Default")
        self.assertEqual(row[CHECK_IN], "08:26")
        self.assertEqual(row[CHECK_OUT], "17:40")
        self.assertEqual(row[ACTUAL], "8:14")
        self.assertEqual(row[REQUIRED], FULL_DAY)
        self.assertEqual(row[LATE], NOTHING)
        self.assertEqual(row[EARLY], NOTHING)
        self.assertEqual(row[ABSENT], NOTHING)

    def test_midday_punches_are_not_paired(self):
        """Four punches on the Tuesday collapse to the ends of the day: only
        the first and the last are the times payroll is given."""
        row = self.day(self.week(), "1001", "01/09/2026")
        self.assertEqual((row[CHECK_IN], row[CHECK_OUT]), ("08:26", "17:40"))

    def test_status_field_is_not_used_to_pair(self):
        """Every punch here carries status=1 ('Check Out') and the days still
        pair correctly — the regression that would silently produce a sheet of
        check-outs with no check-ins."""
        row = self.day(self.week(), "1002", "02/09/2026")
        self.assertEqual((row[CHECK_IN], row[CHECK_OUT]), ("08:20", "12:01"))

    def test_lateness_and_early_departure_are_separate_columns(self):
        late = self.day(self.week(), "1001", "02/09/2026")
        self.assertEqual(late[LATE], "0:13")
        self.assertEqual(late[EARLY], NOTHING)

        early = self.day(self.week(), "1002", "02/09/2026")
        self.assertEqual(early[LATE], NOTHING)
        self.assertEqual(early[EARLY], "4:30")
        self.assertEqual(early[ACTUAL], "3:41")

    def test_a_single_punch_day_is_absent_with_no_check_out(self):
        """Somebody who forgot to punch out is not a zero-hour day. Repeating
        the check-in as the check-out would read as a worked day; the empty
        Check-Out next to a full Vắng mặt is what says the record is broken."""
        row = self.day(self.week(), "1001", "03/09/2026")
        self.assertEqual(row[CHECK_IN], "08:36")
        self.assertIsNone(row[CHECK_OUT])
        self.assertEqual(row[ACTUAL], NOTHING)
        self.assertEqual(row[ABSENT], FULL_DAY)
        self.assertEqual(row[LATE], NOTHING)

    def test_a_day_with_no_punches_is_a_full_absence(self):
        row = self.day(self.week(), "1001", "04/09/2026")
        self.assertIsNone(row[CHECK_IN])
        self.assertIsNone(row[CHECK_OUT])
        self.assertEqual(row[REQUIRED], FULL_DAY)
        self.assertEqual(row[ABSENT], FULL_DAY)

    def test_a_weekend_owes_nothing_and_has_no_timetable(self):
        for day in ("05/09/2026", "06/09/2026"):
            with self.subTest(day=day):
                row = self.day(self.week(), "1001", day)
                self.assertIsNone(row[TIMETABLE])
                self.assertEqual(row[REQUIRED], NOTHING)
                self.assertEqual(row[ABSENT], NOTHING)
                self.assertEqual(row[ACTUAL], NOTHING)

    def test_overtime_columns_are_present_and_empty(self):
        """Kept because the layout is a contract; left at zero because this
        install has no overtime rule to apply, and a payroll sheet must not
        carry hours nobody approved."""
        row = self.day(self.week(), "1001", "01/09/2026")
        self.assertEqual([row[OT1], row[OT2], row[OT3]], [NOTHING] * 3)

    def test_everything_is_written_as_text(self):
        """Durations are elapsed time, not times of day: as real values Excel
        would show 8:14 worked as a quarter past eight in the morning, and 25
        hours of overtime would wrap round to 1:00."""
        ws = self.sheet(self.export(**self.WEEK))
        for coordinate in ("A2", "C2", "E2", "M2"):
            with self.subTest(cell=coordinate):
                self.assertEqual(ws[coordinate].number_format, "@")
        self.assertIsInstance(ws["A2"].value, str)
        self.assertEqual(ws["C2"].value, "01/09/2026")   # dd/mm/yyyy, as text

    def test_the_sheet_looks_like_the_one_it_replaces(self):
        """Tahoma 8 on a grey, centred, bordered header — ZKTime.Net's own
        styling, so the file HR opens next month looks like last month's."""
        ws = self.sheet(self.export(**self.WEEK))
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["A1"].font.name, "Tahoma")
        self.assertEqual(ws["A1"].font.sz, 8)
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "FFD3D3D3")
        self.assertEqual(ws["A1"].alignment.horizontal, "center")
        self.assertEqual(ws["A2"].border.left.style, "thin")

    def test_the_range_comes_from_the_punches_when_the_filter_has_none(self):
        """An unfiltered export is still bounded: 01/09 to 03/09 here, not a
        grid over all of time."""
        rows = self.rows(self.export())
        self.assertEqual([r[DAY] for r in rows[:3]],
                         ["01/09/2026", "02/09/2026", "03/09/2026"])
        self.assertEqual(len(rows), 6)     # three days, two people

    def test_one_employee_filter_gives_one_employee(self):
        rows = self.week(user_id="1001")
        self.assertEqual(len(rows), 6)
        self.assertEqual({r[PIN] for r in rows}, {"1001"})

    def test_a_device_filter_narrows_to_its_enrolled_users(self):
        db = self.Session()
        try:
            db.add(DeviceEmployee(device_sn=OTHER_SN, user_id="1002", uid=1))
            db.commit()
        finally:
            db.close()

        rows = self.week(device_sn=OTHER_SN)
        self.assertEqual({r[PIN] for r in rows}, {"1002"})

    def test_a_device_with_no_enrolment_synced_still_reports_everyone(self):
        """An empty device_employees table means nobody has synced that
        terminal's users, not that nobody is enrolled on it — narrowing by it
        would hand HR an empty sheet."""
        rows = self.week(device_sn=MAIN_SN)
        self.assertEqual({r[PIN] for r in rows}, {"1001", "1002"})

    def test_a_pin_with_no_employee_record_still_gets_rows(self):
        """Somebody punching with a PIN the roster does not know is in the
        sheet, under their PIN, rather than quietly dropped."""
        self.punch(datetime(2026, 9, 2, 9, 0, 0), user_id="9999")
        rows = self.week()
        stranger = self.day(rows, "9999", "02/09/2026")
        self.assertEqual(stranger[NAME], "9999")

    def test_pins_sort_as_numbers(self):
        self.punch(datetime(2026, 9, 2, 9, 0, 0), user_id="211")
        order = []
        for row in self.week():
            if row[PIN] not in order:
                order.append(row[PIN])
        self.assertEqual(order, ["211", "1001", "1002"])

    def test_a_grid_too_large_is_refused_with_the_numbers(self):
        """The ceiling the timesheet actually runs into: a row per person per
        day, built from very few punches, so the punch count says nothing
        about how big the sheet will be."""
        original = config.ATTENDANCE_EXPORT_MAX_ROWS
        config.ATTENDANCE_EXPORT_MAX_ROWS = 11      # the week needs 12
        try:
            res = self.export(**self.WEEK)
        finally:
            config.ATTENDANCE_EXPORT_MAX_ROWS = original

        self.assertEqual(res.status_code, 400)
        detail = res.json()["detail"]
        self.assertIn("12", detail)                 # rows the sheet would have
        self.assertIn("6 days", detail)
        self.assertIn("Narrow the date range", detail)


class ExportGuardTests(ExportTestCase):
    """Refusals, naming, auditing — the parts that are not about cell values."""

    def seed(self):
        self.punch(datetime(2026, 9, 1, 8, 0, 0))
        self.punch(datetime(2026, 9, 1, 17, 0, 0))
        self.punch(datetime(2026, 9, 2, 8, 1, 0), user_id="1002", sn=OTHER_SN, zone=None)

    def test_empty_result_is_an_empty_workbook_not_an_error(self):
        res = self.export(user_id="does-not-exist")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self.rows(res), [])

    def test_filename_carries_the_filter(self):
        res = self.export(device_sn=MAIN_SN,
                          from_date="2026-09-01T00:00:00",
                          to_date="2026-09-30T23:59:59")
        disposition = res.headers["content-disposition"]
        self.assertIn("attendance-timesheet", disposition)
        self.assertIn(MAIN_SN, disposition)
        self.assertIn("20260901", disposition)
        self.assertIn("20260930", disposition)
        # ASCII only — nothing here needs RFC 5987 encoding to survive.
        disposition.encode("ascii")

    def test_too_many_rows_is_refused_with_the_numbers(self):
        """The ceiling counts punches *read*: the sheet is small, but it is
        still built by walking every punch behind it."""
        original = config.ATTENDANCE_EXPORT_MAX_ROWS
        config.ATTENDANCE_EXPORT_MAX_ROWS = 2
        try:
            res = self.export()
        finally:
            config.ATTENDANCE_EXPORT_MAX_ROWS = original

        self.assertEqual(res.status_code, 400)
        detail = res.json()["detail"]
        self.assertIn("3", detail)      # what matched
        self.assertIn("2", detail)      # what is allowed
        self.assertIn("Narrow the date range", detail)

    def test_export_is_audited(self):
        self.export(device_sn=MAIN_SN)

        db = self.Session()
        try:
            entry = db.query(AuditLog).filter_by(action="attendance_export").one()
        finally:
            db.close()

        self.assertEqual(entry.actor, "tester")
        self.assertEqual(entry.ip, "203.0.113.10")
        self.assertIn("rows=2", entry.detail)
        self.assertIn(f"device={MAIN_SN}", entry.detail)

    def test_refused_export_writes_no_audit_row(self):
        original = config.ATTENDANCE_EXPORT_MAX_ROWS
        config.ATTENDANCE_EXPORT_MAX_ROWS = 1
        try:
            self.export()
        finally:
            config.ATTENDANCE_EXPORT_MAX_ROWS = original

        db = self.Session()
        try:
            self.assertEqual(
                db.query(AuditLog).filter_by(action="attendance_export").count(), 0)
        finally:
            db.close()

    def test_a_range_that_ends_before_it_starts_is_refused(self):
        """An inverted range matches nothing, so left alone it would come out
        as an empty timesheet with no error anywhere — a wrong answer that
        looks like a right one. Both endpoints refuse it instead."""
        res = self.export(from_date="2026-09-30T00:00:00",
                          to_date="2026-09-01T00:00:00")
        self.assertEqual(res.status_code, 400)
        self.assertIn("before its start", res.json()["detail"])

        res = self.client.get("/attendance", params={
            "from_date": "2026-09-30T00:00:00",
            "to_date": "2026-09-01T00:00:00",
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn("before its start", res.json()["detail"])

    def test_a_refused_range_writes_no_audit_row(self):
        self.export(from_date="2026-09-30T00:00:00", to_date="2026-09-01T00:00:00")

        db = self.Session()
        try:
            self.assertEqual(
                db.query(AuditLog).filter_by(action="attendance_export").count(), 0)
        finally:
            db.close()

    def test_a_single_instant_range_is_not_inverted(self):
        """From equal to To is a legal, if narrow, filter: both bounds are
        inclusive, so it selects the punches on that second."""
        res = self.export(from_date="2026-09-01T08:00:00",
                          to_date="2026-09-01T08:00:00")
        self.assertEqual(res.status_code, 200)

        res = self.client.get("/attendance", params={
            "from_date": "2026-09-01T08:00:00",
            "to_date": "2026-09-01T08:00:00",
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["total"], 1)

    def test_listing_still_answers_alongside_the_export_route(self):
        """`/attendance/export.xlsx` must not shadow the list endpoint."""
        res = self.client.get("/attendance")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["total"], 3)


if __name__ == "__main__":
    unittest.main()
