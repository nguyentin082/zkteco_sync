"""Rows that belong to nobody: PIN 0, and people since deleted.

    python -m unittest discover -s tests -v

The bug these guard against is not hypothetical. On the live installation the
Attendance screen showed 13,185 records for an "employee" called 0 — every one
of them a failed scan the terminal logged seconds before the same person's
successful punch — plus 38,946 more belonging to PINs no longer on the roster.
52,131 of 96,345 rows named nobody the operator could identify.

Three rules are pinned here, and the third is the one most likely to be
"tidied" away by somebody who sees only the first two:

  1. A PIN-0 record is not attendance and never enters the table, whichever
     protocol carries it.
  2. The screen shows punches belonging to somebody on the roster, and hides
     the rest behind `include_hidden` rather than deleting anything.
  3. The Excel timesheet still contains a deleted employee's punches. Somebody
     who resigned on the 12th worked until the 12th, and that month's sheet is
     reconciled against payroll. Filtering the export the way the screen is
     filtered would take half a month of pay off it silently.
"""

import os
import unittest
from datetime import datetime
from io import BytesIO

# Set before importing anything from `app` — see tests/test_adms.py for why.
os.environ.setdefault("APP_ENV", "development")
os.environ.setdefault("SECRET_KEY", "x" * 48)
os.environ.setdefault("DEFAULT_DEVICE_TIMEZONE", "Asia/Dubai")

from fastapi import FastAPI                                    # noqa: E402
from fastapi.testclient import TestClient                      # noqa: E402
from openpyxl import load_workbook                             # noqa: E402
from sqlalchemy import create_engine                           # noqa: E402
from sqlalchemy.orm import sessionmaker                        # noqa: E402
from sqlalchemy.pool import StaticPool                         # noqa: E402

from app.database import Base, get_db                          # noqa: E402
from app.deps import require_auth                              # noqa: E402
from app.models import AttendanceLog, Device, Employee, User   # noqa: E402
from app.routers import adms as adms_router                    # noqa: E402
from app.routers import attendance as attendance_router        # noqa: E402
from app.services.punch_filter import (                        # noqa: E402
    NON_PERSON_PINS, is_person_pin,
)

SN = "HIDDENDEV0001"

# The status every punch on the operator's live installation carries, because
# nobody presses the mode key on the terminal.
UNPRESSED = 1
# The verify mode every PIN-0 record on the live installation carries. Present
# here so the fixtures look like the real data — deliberately *not* what the
# filter keys on; see `is_person_pin`.
FAILED_SCAN = 201


class PinRuleTests(unittest.TestCase):
    """`is_person_pin`, and the value set that mirrors it into SQL."""

    def test_zero_in_every_spelling_is_nobody(self):
        for pin in ("0", "00", "0000", "", "  ", " 0 ", 0, None):
            self.assertFalse(is_person_pin(pin), pin)

    def test_any_other_pin_is_somebody(self):
        for pin in ("1", "53", "6521592", "A7", 13):
            self.assertTrue(is_person_pin(pin), pin)

    def test_the_sql_set_agrees_with_the_python_rule(self):
        # The two are used in different places on the same question, so they
        # must not be able to disagree about a value the column can hold.
        for pin in NON_PERSON_PINS:
            self.assertFalse(is_person_pin(pin), pin)
        for pin in ("0" * 24, ""):
            self.assertIn(pin, NON_PERSON_PINS)
        for pin in ("1", "53", "0" * 23 + "1"):
            self.assertNotIn(pin, NON_PERSON_PINS)


class HiddenRowsTestCase(unittest.TestCase):
    """One device, one employee on the roster, one since deleted, and the
    failed scans the terminal wrote between them."""

    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(bind=self.engine)
        self.Session = sessionmaker(bind=self.engine, autocommit=False, autoflush=False)

        app = FastAPI()
        app.include_router(attendance_router.router)
        app.include_router(adms_router.router)

        def _override_get_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        operator = User(id=1, username="tester", role="viewer", password_hash="x")
        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[require_auth] = lambda: operator
        self.client = TestClient(app, client=("203.0.113.10", 40000))

        db = self.Session()
        try:
            db.add(Device(serial_number=SN, ip_address="203.0.113.10", port=4370,
                          name="Main Door", status="approved", timezone="Asia/Dubai"))
            # On the roster.
            db.add(Employee(user_id="1001", name="Nguyen Van A"))
            # "1002" is deliberately absent: somebody `delete_employee` removed,
            # whose punches it deliberately left behind.
            db.commit()
        finally:
            db.close()

        # 3 August 2026, a Monday. A punch each for the two people, and the
        # failed scan the terminal logged three seconds before one of them —
        # the shape 90% of the live PIN-0 records have.
        self.punch(datetime(2026, 8, 3, 8, 12, 19), "0", punch=FAILED_SCAN)
        self.punch(datetime(2026, 8, 3, 8, 12, 22), "1001")
        self.punch(datetime(2026, 8, 3, 17, 40, 0), "1001")
        self.punch(datetime(2026, 8, 3, 8, 30, 0), "1002")
        self.punch(datetime(2026, 8, 3, 17, 35, 0), "1002")

    def tearDown(self):
        self.client.close()
        Base.metadata.drop_all(bind=self.engine)
        self.engine.dispose()

    def punch(self, moment, user_id, punch=1):
        db = self.Session()
        try:
            db.add(AttendanceLog(device_sn=SN, user_id=user_id, timestamp=moment,
                                 status=UNPRESSED, punch=punch, source="sdk_pull",
                                 timezone="Asia/Dubai"))
            db.commit()
        finally:
            db.close()

    def listed(self, **params):
        res = self.client.get("/attendance", params=params)
        self.assertEqual(res.status_code, 200, res.text)
        return res.json()


class ScreenTests(HiddenRowsTestCase):
    """GET /attendance — what an operator sees by default."""

    def test_failed_scans_and_deleted_employees_are_both_hidden(self):
        body = self.listed()
        pins = {item["user_id"] for item in body["items"]}
        self.assertEqual(pins, {"1001"})
        self.assertNotIn("0", pins)
        self.assertNotIn("1002", pins)

    def test_the_count_matches_what_is_shown(self):
        # A total of 5 over a table of 2 would give the operator a pager
        # leading to blank pages, and a record count that contradicts the
        # screen. `total` comes off the same filtered query for that reason.
        body = self.listed()
        self.assertEqual(body["total"], 2)
        self.assertEqual(len(body["items"]), 2)

    def test_include_hidden_brings_every_row_back(self):
        body = self.listed(include_hidden="true")
        pins = sorted(item["user_id"] for item in body["items"])
        self.assertEqual(pins, ["0", "1001", "1001", "1002", "1002"])
        self.assertEqual(body["total"], 5)

    def test_nothing_was_deleted_to_achieve_this(self):
        # The whole design rests on the rows still being there: this is a
        # default, not an erasure, and the flag above is the way back.
        db = self.Session()
        try:
            self.assertEqual(db.query(AttendanceLog).count(), 5)
        finally:
            db.close()

    def test_hiding_survives_a_narrower_filter(self):
        body = self.listed(device_sn=SN, from_date="2026-08-03T00:00:00",
                           to_date="2026-08-03T23:59:59")
        self.assertEqual({i["user_id"] for i in body["items"]}, {"1001"})

    def test_a_deleted_employee_can_still_be_looked_up_by_pin(self):
        # Naming a PIN is a deliberate lookup, so the flag has to be the way
        # to answer it. Without this the operator gets "No records found" for
        # somebody whose punches are demonstrably in the database.
        self.assertEqual(self.listed(user_id="1002")["total"], 0)
        self.assertEqual(
            self.listed(user_id="1002", include_hidden="true")["total"], 2
        )


class TimesheetTests(HiddenRowsTestCase):
    """GET /attendance/export.xlsx — filtered differently, on purpose."""

    def sheet_rows(self):
        res = self.client.get("/attendance/export.xlsx", params={
            "from_date": "2026-08-03T00:00:00",
            "to_date": "2026-08-03T23:59:59",
        })
        self.assertEqual(res.status_code, 200, res.text)
        ws = load_workbook(BytesIO(res.content)).worksheets[0]
        return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

    def test_a_deleted_employee_keeps_their_hours(self):
        # The point of the whole asymmetry. "1002" has no roster row and still
        # gets their day, with both punches on it.
        by_pin = {row[0]: row for row in self.sheet_rows()}
        self.assertIn("1002", by_pin)
        self.assertEqual(by_pin["1002"][12], "08:30")   # first punch
        self.assertEqual(by_pin["1002"][13], "17:35")   # last punch

    def test_no_phantom_person_called_zero(self):
        # `_report_people` unions everybody present in the records, so before
        # the filter this produced a full timesheet row for "0" — with hours
        # scored off failed scans.
        self.assertNotIn("0", {row[0] for row in self.sheet_rows()})

    def test_a_failed_scan_does_not_become_somebody_else_s_arrival(self):
        # The 08:12:19 failed scan sits three seconds before 1001's real
        # 08:12:22 punch. It carries its own PIN, so it could never have been
        # attributed to them — this pins that down, since the day is paired by
        # earliest and latest punch and an off-by-one in the filter would show
        # up here first.
        by_pin = {row[0]: row for row in self.sheet_rows()}
        self.assertEqual(by_pin["1001"][12], "08:12")
        self.assertEqual(by_pin["1001"][13], "17:40")


class IngestTests(unittest.TestCase):
    """Neither writer may put a PIN-0 record in the table again.

    The ADMS push path is exercised over HTTP. The SDK pull path — the one
    that actually let these in, and the only one this installation uses — is
    exercised at its filter, since the rest of `pull_attendance` is a live
    TCP session with a terminal.
    """

    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(bind=self.engine)
        self.Session = sessionmaker(bind=self.engine, autocommit=False, autoflush=False)

        app = FastAPI()
        app.include_router(adms_router.router)

        def _override_get_db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = _override_get_db
        self.client = TestClient(app, client=("203.0.113.10", 40000))

        db = self.Session()
        try:
            db.add(Device(serial_number=SN, ip_address="203.0.113.10", port=4370,
                          status="approved", timezone="Asia/Dubai"))
            db.commit()
        finally:
            db.close()

    def tearDown(self):
        self.client.close()
        Base.metadata.drop_all(bind=self.engine)
        self.engine.dispose()

    def test_an_attlog_push_drops_pin_zero_and_keeps_the_rest(self):
        body = (
            "0\t2026-08-03 08:12:19\t1\t201\n"
            "1001\t2026-08-03 08:12:22\t1\t1\n"
        )
        res = self.client.post(
            f"/iclock/cdata?SN={SN}&table=ATTLOG", content=body.encode(),
        )
        self.assertEqual(res.status_code, 200)
        # The device must still be told the upload was accepted — a terminal
        # that does not see its acknowledgement retries forever.
        self.assertTrue(res.text.strip().startswith("OK"))

        db = self.Session()
        try:
            pins = [r.user_id for r in db.query(AttendanceLog).all()]
        finally:
            db.close()
        self.assertEqual(pins, ["1001"])

    def test_the_sdk_pull_filter_is_the_same_rule(self):
        # pyzk hands back `att.user_id` as the device wrote it; the pull loop
        # asks exactly this before storing anything.
        from app.services import poller
        self.assertIs(poller.is_person_pin, is_person_pin)
        self.assertFalse(poller.is_person_pin("0"))
        self.assertTrue(poller.is_person_pin("1001"))


if __name__ == "__main__":
    unittest.main()
